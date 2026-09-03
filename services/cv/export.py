"""
Export d'un CV ciblé en DOCX et en PDF.

Gabarit calqué sur le CV existant de l'utilisateur (fourni comme
référence) : en-tête compact, compétences regroupées par catégorie,
expériences en deux colonnes (dates/lieu à gauche, contenu à droite),
formation et certifications, langues et centres d'intérêt.

L'export reste une mise en forme : il n'ajoute, ne reformule et ne
complète aucun contenu. Tout ce qui apparaît dans le fichier produit
vient de l'objet TargetedCV, donc du Master CV.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date
from pathlib import Path

from services.cv.results import TargetedCV, CVEducation, CVCertification


EXPORT_DIR = (
    Path(__file__).resolve().parent.parent.parent / "exports"
)


# Sections activables/désactivables à l'export. `included_sections`
# vaut None par défaut (tout est inclus) : c'est le candidat qui
# choisit, pour une offre donnée, ce qui reste pertinent à montrer —
# jamais le contenu lui-même, seulement sa présence.
ALL_CV_SECTIONS = frozenset(
    {
        "resume",
        "competences",
        "experiences",
        "formation_certifications",
        "langues",
        "interets",
    }
)


def _inclut(section: str, included_sections: set[str] | None) -> bool:
    return included_sections is None or section in included_sections


MOIS_ABREGES = (
    "Janv.", "Févr.", "Mars", "Avr.", "Mai", "Juin",
    "Juil.", "Août", "Sept.", "Oct.", "Nov.", "Déc.",
)


def _format_periode(
    debut: date,
    fin: date | None,
) -> str:

    def _mois_annee(valeur: date) -> str:
        return f"{MOIS_ABREGES[valeur.month - 1]} {valeur.year}"

    if fin is None:
        return f"{_mois_annee(debut)} – Aujourd'hui"

    return f"{_mois_annee(debut)} – {_mois_annee(fin)}"


def _slugify(value: str) -> str:

    normalise = unicodedata.normalize("NFKD", value)

    normalise = "".join(
        caractere
        for caractere in normalise
        if not unicodedata.combining(caractere)
    )

    normalise = re.sub(r"[^a-zA-Z0-9]+", "-", normalise)

    return normalise.strip("-").lower()[:60] or "cv"


def default_export_path(
    cv: TargetedCV,
    extension: str,
) -> Path:
    """
    Chemin de sortie par défaut, stable pour un même CV et une même
    offre : régénérer écrase le fichier au lieu d'en empiler.
    """

    nom = _slugify(cv.full_name)
    poste = _slugify(cv.job_offer_title)

    return EXPORT_DIR / f"cv-{nom}-{poste}.{extension}"


# ============================================================
# FORMATION + CERTIFICATIONS — LIGNE COMMUNE TRIEE
# ============================================================

def _education_line(item: CVEducation) -> tuple[str, str]:

    intitule = item.degree

    if item.field_of_study:
        intitule += f", {item.field_of_study}"

    if item.start_year and item.end_year:
        annees = f"{item.start_year} – {item.end_year}"
    elif item.end_year:
        annees = str(item.end_year)
    elif item.start_year:
        annees = str(item.start_year)
    else:
        annees = ""

    return annees, f"{intitule} — {item.institution}"


def _certification_line(item: CVCertification) -> tuple[str, str]:

    annees = str(item.obtained_year) if item.obtained_year else ""

    intitule = item.name

    if item.organization:
        intitule += f" — {item.organization}"

    return annees, intitule


def _formation_certifications_lines(
    cv: TargetedCV,
) -> list[tuple[str, str]]:
    """
    Fusionne formation et certifications en une seule liste, triée de
    la plus récente à la plus ancienne.
    """

    lignes = [
        (item.end_year or item.start_year or 0, _education_line(item))
        for item in cv.educations
    ] + [
        (item.obtained_year or 0, _certification_line(item))
        for item in cv.certifications
    ]

    lignes.sort(key=lambda paire: paire[0], reverse=True)

    return [ligne for _annee, ligne in lignes]


# ============================================================
# DOCX
# ============================================================

def export_docx(
    cv: TargetedCV,
    path: Path | str | None = None,
    included_sections: set[str] | None = None,
) -> Path:
    """
    Écrit le CV ciblé dans un fichier .docx et retourne son chemin.

    `included_sections` restreint les rubriques affichées (voir
    ALL_CV_SECTIONS) ; None (par défaut) inclut tout ce que le
    contenu du CV rend pertinent, comme avant l'ajout de ce paramètre.
    """

    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    destination = Path(
        path
        if path is not None
        else default_export_path(cv, "docx")
    )

    destination.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    for section in document.sections:
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)

    GRIS = RGBColor(0x55, 0x55, 0x55)

    def _rubrique(titre: str) -> None:
        paragraphe = document.add_paragraph()
        run = paragraphe.add_run(titre.upper())
        run.bold = True
        run.font.size = Pt(12)
        paragraphe.paragraph_format.space_before = Pt(12)
        paragraphe.paragraph_format.space_after = Pt(4)

    # --------------------------------------------------------
    # EN-TETE
    # --------------------------------------------------------

    nom = document.add_paragraph()
    run_nom = nom.add_run(cv.full_name.upper())
    run_nom.bold = True
    run_nom.font.size = Pt(20)
    nom.paragraph_format.space_after = Pt(2)

    titre = cv.cv_title or cv.headline

    if titre:
        accroche = document.add_paragraph()
        run_accroche = accroche.add_run(titre)
        run_accroche.bold = True
        run_accroche.font.size = Pt(12)
        accroche.paragraph_format.space_after = Pt(4)

    coordonnees = " | ".join(
        partie
        for partie in (
            cv.availability,
            cv.location,
            cv.phone,
            cv.email,
            cv.linkedin_url,
        )
        if partie
    )

    if coordonnees:
        ligne_contact = document.add_paragraph()
        run_contact = ligne_contact.add_run(coordonnees)
        run_contact.font.size = Pt(9)
        run_contact.font.color.rgb = GRIS
        ligne_contact.paragraph_format.space_after = Pt(4)

    # --------------------------------------------------------
    # PROFIL
    # --------------------------------------------------------

    if cv.summary and _inclut("resume", included_sections):
        _rubrique("Profil")
        document.add_paragraph(cv.summary)

    # --------------------------------------------------------
    # COMPETENCES CLES
    # --------------------------------------------------------
    #
    # Uniquement les compétences prouvées, regroupées par catégorie
    # du référentiel : c'est la garantie d'honnêteté du CV généré.

    if cv.skill_groups and _inclut("competences", included_sections):

        _rubrique("Compétences clés")

        for groupe in cv.skill_groups:

            ligne = document.add_paragraph(style="List Bullet")

            run_categorie = ligne.add_run(f"{groupe.category} : ")
            run_categorie.bold = True

            ligne.add_run(", ".join(groupe.skills))

    # --------------------------------------------------------
    # EXPERIENCES — DEUX COLONNES (DATES/LIEU | CONTENU)
    # --------------------------------------------------------

    if cv.experiences and _inclut("experiences", included_sections):

        _rubrique("Expériences professionnelles")

        for experience in cv.experiences:

            table = document.add_table(rows=1, cols=2)
            table.autofit = False

            colonne_date, colonne_contenu = table.rows[0].cells

            colonne_date.width = Cm(3.6)
            colonne_contenu.width = Cm(12.8)

            p_periode = colonne_date.paragraphs[0]
            run_periode = p_periode.add_run(
                _format_periode(
                    experience.start_date,
                    experience.end_date,
                )
            )
            run_periode.font.size = Pt(9)
            run_periode.font.color.rgb = GRIS

            if experience.location:
                p_lieu = colonne_date.add_paragraph()
                run_lieu = p_lieu.add_run(experience.location)
                run_lieu.font.size = Pt(9)
                run_lieu.font.color.rgb = GRIS

            p_titre = colonne_contenu.paragraphs[0]
            run_titre = p_titre.add_run(
                f"{experience.job_title} — {experience.company}"
            )
            run_titre.bold = True

            # Le contexte d'entreprise situe la mission : sans lui,
            # une puce comme « pilotage de projets de bout en bout »
            # ne dit ni sur quoi, ni dans quel environnement.
            if experience.business_context:

                p_contexte = colonne_contenu.add_paragraph()
                run_contexte = p_contexte.add_run(
                    experience.business_context
                )
                run_contexte.italic = True
                run_contexte.font.size = Pt(9)
                run_contexte.font.color.rgb = GRIS
                p_contexte.paragraph_format.space_after = Pt(2)

            # Les réalisations d'abord : c'est ce qui distingue le
            # candidat, et le seul endroit du Master CV qui porte des
            # chiffres.
            for realisation in experience.achievement_lines:

                puce = colonne_contenu.add_paragraph(
                    style="List Bullet"
                )

                run_titre_realisation = puce.add_run(
                    realisation.title
                )
                run_titre_realisation.bold = True

                if realisation.detail:
                    puce.add_run(f" — {realisation.detail}")

            for ligne in experience.lines:
                colonne_contenu.add_paragraph(
                    ligne.text,
                    style="List Bullet",
                )

            # Espace visuel entre deux expériences.
            document.add_paragraph().paragraph_format.space_after = (
                Pt(2)
            )

    # --------------------------------------------------------
    # FORMATION & CERTIFICATIONS
    # --------------------------------------------------------

    lignes_formation = _formation_certifications_lines(cv)

    if lignes_formation and _inclut(
        "formation_certifications", included_sections
    ):

        _rubrique("Formation & certifications")

        for annee, intitule in lignes_formation:

            table = document.add_table(rows=1, cols=2)
            table.autofit = False

            colonne_annee, colonne_intitule = table.rows[0].cells

            colonne_annee.width = Cm(3.6)
            colonne_intitule.width = Cm(12.8)

            run_annee = colonne_annee.paragraphs[0].add_run(annee)
            run_annee.font.size = Pt(9)
            run_annee.font.color.rgb = GRIS

            colonne_intitule.paragraphs[0].add_run(intitule)

    # --------------------------------------------------------
    # LANGUES & CENTRES D'INTERET
    # --------------------------------------------------------

    afficher_langues = bool(cv.languages) and _inclut(
        "langues", included_sections
    )
    afficher_interets = bool(cv.interests) and _inclut(
        "interets", included_sections
    )

    if afficher_langues or afficher_interets:

        _rubrique("Langues & centres d'intérêt")

        if afficher_langues:
            document.add_paragraph(f"Langues : {cv.languages}")

        if afficher_interets:
            document.add_paragraph(
                f"Centres d'intérêt : {cv.interests}"
            )

    document.save(str(destination))

    return destination


# ============================================================
# PDF
# ============================================================

def export_pdf(
    cv: TargetedCV,
    path: Path | str | None = None,
    included_sections: set[str] | None = None,
) -> Path:
    """
    Écrit le CV ciblé dans un fichier .pdf et retourne son chemin.

    `included_sections` restreint les rubriques affichées (voir
    ALL_CV_SECTIONS) ; None (par défaut) inclut tout ce que le
    contenu du CV rend pertinent, comme avant l'ajout de ce paramètre.
    """

    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import (
        ParagraphStyle,
        getSampleStyleSheet,
    )
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    destination = Path(
        path
        if path is not None
        else default_export_path(cv, "pdf")
    )

    destination.parent.mkdir(parents=True, exist_ok=True)

    styles = getSampleStyleSheet()

    GRIS = HexColor("#555555")

    style_nom = ParagraphStyle(
        "NomCandidat",
        parent=styles["Title"],
        alignment=0,
        fontSize=20,
        leading=24,
        spaceAfter=2,
    )

    style_accroche = ParagraphStyle(
        "Accroche",
        parent=styles["Normal"],
        fontSize=11,
        leading=14,
        spaceAfter=4,
    )

    style_contact = ParagraphStyle(
        "Contact",
        parent=styles["Normal"],
        fontSize=9,
        textColor=GRIS,
        spaceAfter=6,
    )

    style_rubrique = ParagraphStyle(
        "Rubrique",
        parent=styles["Heading2"],
        fontSize=12,
        spaceBefore=12,
        spaceAfter=4,
    )

    style_normal = styles["Normal"]

    style_petit = ParagraphStyle(
        "Petit",
        parent=styles["Normal"],
        fontSize=9,
        textColor=GRIS,
    )

    style_titre_poste = ParagraphStyle(
        "TitrePoste",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
    )

    document = SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title=f"CV {cv.full_name}",
    )

    elements = [Paragraph(cv.full_name.upper(), style_nom)]

    titre = cv.cv_title or cv.headline

    if titre:
        elements.append(Paragraph(titre, style_accroche))

    coordonnees = " | ".join(
        partie
        for partie in (
            cv.availability,
            cv.location,
            cv.phone,
            cv.email,
            cv.linkedin_url,
        )
        if partie
    )

    if coordonnees:
        elements.append(Paragraph(coordonnees, style_contact))

    # --------------------------------------------------------
    # PROFIL
    # --------------------------------------------------------

    if cv.summary and _inclut("resume", included_sections):
        elements.append(Paragraph("PROFIL", style_rubrique))
        elements.append(Paragraph(cv.summary, style_normal))

    # --------------------------------------------------------
    # COMPETENCES CLES
    # --------------------------------------------------------

    if cv.skill_groups and _inclut("competences", included_sections):

        elements.append(
            Paragraph("COMPÉTENCES CLÉS", style_rubrique)
        )

        elements.append(
            ListFlowable(
                [
                    ListItem(
                        Paragraph(
                            f"<b>{groupe.category} :</b> "
                            + ", ".join(groupe.skills),
                            style_normal,
                        )
                    )
                    for groupe in cv.skill_groups
                ],
                bulletType="bullet",
                leftIndent=12,
            )
        )

    # --------------------------------------------------------
    # EXPERIENCES — TABLEAU 2 COLONNES PAR EXPERIENCE
    # --------------------------------------------------------

    if cv.experiences and _inclut("experiences", included_sections):

        elements.append(
            Paragraph(
                "EXPÉRIENCES PROFESSIONNELLES", style_rubrique
            )
        )

        for experience in cv.experiences:

            colonne_date = [
                Paragraph(
                    _format_periode(
                        experience.start_date,
                        experience.end_date,
                    ),
                    style_petit,
                )
            ]

            if experience.location:
                colonne_date.append(
                    Paragraph(experience.location, style_petit)
                )

            colonne_contenu = [
                Paragraph(
                    f"{experience.job_title} — {experience.company}",
                    style_titre_poste,
                )
            ]

            # Le contexte d'entreprise situe la mission : sans lui,
            # une puce comme « pilotage de projets de bout en bout »
            # ne dit ni sur quoi, ni dans quel environnement.
            if experience.business_context:
                colonne_contenu.append(
                    Paragraph(
                        f"<i>{experience.business_context}</i>",
                        style_petit,
                    )
                )

            # Les réalisations d'abord : c'est ce qui distingue le
            # candidat, et le seul endroit du Master CV qui porte des
            # chiffres.
            puces = [
                Paragraph(
                    f"<b>{realisation.title}</b>"
                    + (
                        f" — {realisation.detail}"
                        if realisation.detail
                        else ""
                    ),
                    style_normal,
                )
                for realisation in experience.achievement_lines
            ] + [
                Paragraph(ligne.text, style_normal)
                for ligne in experience.lines
            ]

            if puces:
                colonne_contenu.append(
                    ListFlowable(
                        [ListItem(puce) for puce in puces],
                        bulletType="bullet",
                        leftIndent=10,
                    )
                )

            table = Table(
                [[colonne_date, colonne_contenu]],
                colWidths=[3.6 * cm, 12.8 * cm],
            )

            table.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )

            elements.append(table)

    # --------------------------------------------------------
    # FORMATION & CERTIFICATIONS
    # --------------------------------------------------------

    lignes_formation = _formation_certifications_lines(cv)

    if lignes_formation and _inclut(
        "formation_certifications", included_sections
    ):

        elements.append(
            Paragraph("FORMATION & CERTIFICATIONS", style_rubrique)
        )

        table = Table(
            [
                [
                    Paragraph(annee, style_petit),
                    Paragraph(intitule, style_normal),
                ]
                for annee, intitule in lignes_formation
            ],
            colWidths=[3.6 * cm, 12.8 * cm],
        )

        table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )

        elements.append(table)

    # --------------------------------------------------------
    # LANGUES & CENTRES D'INTERET
    # --------------------------------------------------------

    afficher_langues = bool(cv.languages) and _inclut(
        "langues", included_sections
    )
    afficher_interets = bool(cv.interests) and _inclut(
        "interets", included_sections
    )

    if afficher_langues or afficher_interets:

        elements.append(
            Paragraph(
                "LANGUES & CENTRES D'INTÉRÊT", style_rubrique
            )
        )

        if afficher_langues:
            elements.append(
                Paragraph(f"Langues : {cv.languages}", style_normal)
            )

        if afficher_interets:
            elements.append(
                Paragraph(
                    f"Centres d'intérêt : {cv.interests}",
                    style_normal,
                )
            )

    document.build(elements)

    return destination
