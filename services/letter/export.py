"""
Export d'une lettre de motivation en DOCX et en PDF.

Comme pour le CV, l'export est une mise en forme : il n'ajoute ni ne
reformule aucun contenu.
"""

from __future__ import annotations

from pathlib import Path

from services.cv.export import EXPORT_DIR, _slugify
from services.letter.results import CoverLetter


# Noms de mois complets pour une date en prose ("2 septembre 2026") :
# distinct des abréviations utilisées sur le CV ("Sept. 2026"), le
# registre d'une lettre est plus formel.
MOIS_COMPLETS = (
    "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre",
    "décembre",
)


def _date_en_toutes_lettres(valeur) -> str:
    return (
        f"{valeur.day} {MOIS_COMPLETS[valeur.month - 1]} {valeur.year}"
    )


def default_letter_path(
    letter: CoverLetter,
    extension: str,
) -> Path:
    """Chemin de sortie stable pour un couple candidat / offre."""

    nom = _slugify(letter.full_name)
    poste = _slugify(letter.job_offer_title)

    return EXPORT_DIR / f"lettre-{nom}-{poste}.{extension}"


# ============================================================
# DOCX
# ============================================================

def export_letter_docx(
    letter: CoverLetter,
    path: Path | str | None = None,
) -> Path:

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    destination = Path(
        path
        if path is not None
        else default_letter_path(letter, "docx")
    )

    destination.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    # --------------------------------------------------------
    # EN-TETE
    # --------------------------------------------------------

    document.add_paragraph(letter.full_name)

    for ligne in (letter.location, letter.phone, letter.email):
        if ligne:
            document.add_paragraph(ligne)

    if letter.company:
        document.add_paragraph("")
        document.add_paragraph(letter.company)

    lieu_date = document.add_paragraph(
        (
            f"{letter.location}, le "
            if letter.location
            else "Le "
        )
        + _date_en_toutes_lettres(letter.redaction_date)
    )
    lieu_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # --------------------------------------------------------
    # OBJET
    # --------------------------------------------------------

    objet = document.add_paragraph()
    objet.add_run(letter.objet).bold = True

    # --------------------------------------------------------
    # CORPS
    # --------------------------------------------------------

    document.add_paragraph(letter.salutation)

    for paragraph in letter.paragraphs:
        document.add_paragraph(paragraph.text)

    document.add_paragraph(letter.closing)

    document.add_paragraph("")
    document.add_paragraph(letter.signature)

    document.save(str(destination))

    return destination


# ============================================================
# PDF
# ============================================================

def export_letter_pdf(
    letter: CoverLetter,
    path: Path | str | None = None,
) -> Path:

    from reportlab.lib.enums import TA_JUSTIFY, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import (
        ParagraphStyle,
        getSampleStyleSheet,
    )
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    destination = Path(
        path
        if path is not None
        else default_letter_path(letter, "pdf")
    )

    destination.parent.mkdir(parents=True, exist_ok=True)

    styles = getSampleStyleSheet()

    style_corps = ParagraphStyle(
        "Corps",
        parent=styles["Normal"],
        alignment=TA_JUSTIFY,
        spaceAfter=10,
        leading=15,
    )

    style_droite = ParagraphStyle(
        "Droite",
        parent=styles["Normal"],
        alignment=TA_RIGHT,
        spaceAfter=16,
    )

    document = SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=f"Lettre de motivation — {letter.full_name}",
    )

    elements = [Paragraph(letter.full_name, styles["Normal"])]

    for ligne in (letter.location, letter.phone, letter.email):
        if ligne:
            elements.append(
                Paragraph(ligne, styles["Normal"])
            )

    if letter.company:
        elements.append(Spacer(1, 12))
        elements.append(
            Paragraph(letter.company, styles["Normal"])
        )

    elements.append(
        Paragraph(
            (
                f"{letter.location}, le "
                if letter.location
                else "Le "
            )
            + _date_en_toutes_lettres(letter.redaction_date),
            style_droite,
        )
    )

    elements.append(
        Paragraph(f"<b>{letter.objet}</b>", style_corps)
    )

    elements.append(
        Paragraph(letter.salutation, style_corps)
    )

    for paragraph in letter.paragraphs:
        elements.append(
            Paragraph(paragraph.text, style_corps)
        )

    elements.append(Paragraph(letter.closing, style_corps))

    elements.append(Spacer(1, 16))

    elements.append(
        Paragraph(letter.signature, styles["Normal"])
    )

    document.build(elements)

    return destination
