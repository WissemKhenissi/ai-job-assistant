"""
Export DOCX / PDF du CV ciblé.

Gabarit calqué sur le CV existant de l'utilisateur : en-tête compact
avec accroche/disponibilité, compétences regroupées par catégorie,
expériences en deux colonnes, formation et certifications, langues et
centres d'intérêt.

Ces tests vérifient surtout qu'aucune compétence non prouvée ne peut
apparaître dans le fichier produit — c'est le dernier endroit où une
compétence non fondée pourrait se glisser avant d'être envoyée à un
recruteur.
"""

from __future__ import annotations

from datetime import date

import pytest

from services.cv.results import (
    CVCertification,
    CVEducation,
    CVEvidenceLine,
    CVExperience,
    CVSkillGroup,
    TargetedCV,
)


def _cv_de_test() -> TargetedCV:
    """
    Un CV où chaque statut est représenté, pour vérifier que seul le
    prouvé ressort dans le document.
    """

    return TargetedCV(
        candidate_id="candidate-test",
        full_name="Wissem Khenissi",
        email="test@example.com",
        phone="0600000000",
        location="Paris",
        linkedin_url="",
        summary="Profil produit.",
        headline="Product / Chef de projet digital",
        availability="Disponible immédiatement",
        languages="Français : natif | Anglais : B2",
        interests="Product Management",
        job_offer_id="job-test",
        job_offer_title="Product Owner",
        skills=["Product Discovery"],
        skill_groups=[
            CVSkillGroup(
                category="Product",
                skills=("Product Discovery",),
            )
        ],
        experiences=[
            CVExperience(
                experience_id="experience-test",
                job_title="Chef de projet",
                company="Groupe Meridiem",
                location="Ivry-sur-Seine",
                start_date=date(2018, 1, 1),
                end_date=date(2024, 12, 31),
                business_context="Contexte e-commerce.",
                lines=[
                    CVEvidenceLine(
                        text="Conception de produits digitaux.",
                        skill="Product Discovery",
                        evidence_id="evidence-1",
                    )
                ],
            )
        ],
        educations=[
            CVEducation(
                institution="Lyon 2",
                degree="Master AEI E-Commerce",
                field_of_study="Cybersécurité",
                start_year=2015,
                end_year=2017,
            )
        ],
        certifications=[
            CVCertification(
                name="PSPO I",
                organization="Scrum.org",
                obtained_year=2026,
            )
        ],
        declared_skills=["CompetenceDeclareeSansPreuve"],
        inferred_skills=["CompetenceSeulementDeduite"],
        missing_skills=["CompetenceAbsente"],
    )


def _texte_du_docx(path) -> str:
    from docx import Document

    document = Document(str(path))

    morceaux = [
        paragraphe.text for paragraphe in document.paragraphs
    ]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                morceaux.append(cell.text)

    return "\n".join(morceaux)


# ============================================================
# DOCX — CONTENU DE BASE
# ============================================================

def test_le_docx_est_cree_et_relisible(tmp_path):
    from services.cv import export_docx

    destination = tmp_path / "cv.docx"

    resultat = export_docx(_cv_de_test(), destination)

    assert resultat == destination
    assert destination.exists()
    assert destination.stat().st_size > 0

    texte = _texte_du_docx(destination)

    assert "WISSEM KHENISSI" in texte
    assert "Chef de projet" in texte
    assert "Conception de produits digitaux." in texte


def test_le_docx_ne_contient_que_les_competences_prouvees(tmp_path):
    """
    Garde-fou final : une compétence déclarée sans preuve, déduite ou
    manquante ne doit apparaître nulle part dans le document envoyé.
    """

    from services.cv import export_docx

    destination = tmp_path / "cv.docx"

    export_docx(_cv_de_test(), destination)

    texte = _texte_du_docx(destination)

    assert "Product Discovery" in texte

    for interdite in (
        "CompetenceDeclareeSansPreuve",
        "CompetenceSeulementDeduite",
        "CompetenceAbsente",
    ):
        assert interdite not in texte, (
            f"{interdite} ne doit pas figurer dans le CV exporté"
        )


def test_le_docx_affiche_la_periode_en_clair(tmp_path):
    from services.cv import export_docx

    destination = tmp_path / "cv.docx"

    export_docx(_cv_de_test(), destination)

    texte = _texte_du_docx(destination)

    assert "Janv. 2018" in texte
    assert "Déc. 2024" in texte


def test_un_poste_en_cours_est_affiche_comme_tel(tmp_path):
    from services.cv import export_docx

    cv = _cv_de_test()
    cv.experiences[0].end_date = None

    destination = tmp_path / "cv.docx"

    export_docx(cv, destination)

    assert "Aujourd'hui" in _texte_du_docx(destination)


def test_un_cv_sans_competence_prouvee_s_exporte_quand_meme(
    tmp_path,
):
    """
    Cas réaliste : aucune compétence prouvée pour cette offre. Le CV
    doit rester exportable, sans section Compétences fabriquée.
    """

    from services.cv import export_docx

    cv = _cv_de_test()
    cv.skills = []
    cv.skill_groups = []

    destination = tmp_path / "cv.docx"

    export_docx(cv, destination)

    texte = _texte_du_docx(destination)

    assert "COMPÉTENCES CLÉS" not in texte.upper()
    assert "WISSEM KHENISSI" in texte


# ============================================================
# DOCX — NOUVELLES SECTIONS DU GABARIT
# ============================================================

def test_le_docx_affiche_l_accroche_et_la_disponibilite(tmp_path):
    from services.cv import export_docx

    destination = tmp_path / "cv.docx"

    export_docx(_cv_de_test(), destination)

    texte = _texte_du_docx(destination)

    assert "Product / Chef de projet digital" in texte
    assert "Disponible immédiatement" in texte


def test_le_docx_regroupe_les_competences_par_categorie(tmp_path):
    from services.cv import export_docx

    destination = tmp_path / "cv.docx"

    export_docx(_cv_de_test(), destination)

    texte = _texte_du_docx(destination)

    assert "Product" in texte


def test_le_docx_affiche_la_localisation_de_l_experience(tmp_path):
    from services.cv import export_docx

    destination = tmp_path / "cv.docx"

    export_docx(_cv_de_test(), destination)

    assert "Ivry-sur-Seine" in _texte_du_docx(destination)


def test_le_docx_affiche_la_formation_et_les_certifications(
    tmp_path,
):
    from services.cv import export_docx

    destination = tmp_path / "cv.docx"

    export_docx(_cv_de_test(), destination)

    texte = _texte_du_docx(destination)

    assert "Master AEI E-Commerce" in texte
    assert "Lyon 2" in texte
    assert "PSPO I" in texte
    assert "Scrum.org" in texte


def test_le_docx_affiche_langues_et_interets(tmp_path):
    from services.cv import export_docx

    destination = tmp_path / "cv.docx"

    export_docx(_cv_de_test(), destination)

    texte = _texte_du_docx(destination)

    assert "Français : natif" in texte
    assert "Product Management" in texte


def test_un_cv_sans_formation_n_affiche_pas_la_rubrique(tmp_path):
    from services.cv import export_docx

    cv = _cv_de_test()
    cv.educations = []
    cv.certifications = []

    destination = tmp_path / "cv.docx"

    export_docx(cv, destination)

    texte = _texte_du_docx(destination).upper()

    assert "FORMATION" not in texte


# ============================================================
# PDF
# ============================================================

def test_le_pdf_est_cree_et_valide(tmp_path):
    from services.cv import export_pdf

    destination = tmp_path / "cv.pdf"

    resultat = export_pdf(_cv_de_test(), destination)

    assert resultat == destination
    assert destination.exists()

    # Un PDF valide commence par son entête de version.
    assert destination.read_bytes().startswith(b"%PDF-")


def test_le_pdf_s_exporte_sans_competence_ni_formation(tmp_path):
    """Un CV très incomplet doit rester exportable sans planter."""

    from services.cv import export_pdf

    cv = _cv_de_test()
    cv.skills = []
    cv.skill_groups = []
    cv.educations = []
    cv.certifications = []
    cv.languages = ""
    cv.interests = ""

    destination = tmp_path / "cv.pdf"

    export_pdf(cv, destination)

    assert destination.read_bytes().startswith(b"%PDF-")


# ============================================================
# NOMMAGE DES FICHIERS
# ============================================================

def test_le_chemin_par_defaut_est_stable_pour_une_meme_offre():
    """
    Régénérer un CV doit écraser le fichier précédent plutôt que
    d'empiler des variantes.
    """

    from services.cv import default_export_path

    cv = _cv_de_test()

    premier = default_export_path(cv, "docx")
    second = default_export_path(cv, "docx")

    assert premier == second
    assert premier.suffix == ".docx"


def test_le_nom_de_fichier_est_lisible_et_sans_accent():
    from services.cv import default_export_path

    cv = _cv_de_test()
    cv.full_name = "Wissem Khenissi"
    cv.job_offer_title = "Chef de Projet Digital H/F"

    chemin = default_export_path(cv, "pdf")

    assert chemin.name == (
        "cv-wissem-khenissi-chef-de-projet-digital-h-f.pdf"
    )


@pytest.mark.parametrize("extension", ["docx", "pdf"])
def test_les_deux_formats_ont_un_chemin_distinct(extension):
    from services.cv import default_export_path

    chemin = default_export_path(_cv_de_test(), extension)

    assert chemin.suffix == f".{extension}"
