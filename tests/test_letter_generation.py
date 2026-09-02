"""
Génération de la lettre de motivation.

La lettre est le document le plus exposé au risque d'affirmation
excessive : c'est un texte suivi, où une formule un peu enthousiaste
peut transformer une hypothèse en fait. Ces tests vérifient qu'elle
n'affirme que ce que le Master CV documente.
"""

from __future__ import annotations

from datetime import date

from conftest import (
    add_candidate,
    add_candidate_skill,
    add_catalog_skill,
    add_evidence,
)


CANDIDATE_ID = "candidate-test"
JOB_OFFER_ID = "job-test"
EXPERIENCE_ID = "experience-test"


def _prepare(
    session,
    competence: str = "Product Discovery",
    avec_preuve: bool = True,
    company: str = "Groupe Meridiem",
):
    from database.models import EvidenceDB, ExperienceDB
    from models.job import JobOfferDB

    add_candidate(session)

    session.add(
        ExperienceDB(
            id=EXPERIENCE_ID,
            candidate_id=CANDIDATE_ID,
            company="Groupe Meridiem",
            job_title="Chef de projet",
            start_date=date(2018, 1, 1),
            end_date=date(2024, 12, 31),
            description="",
            business_context="",
            team_context="",
        )
    )

    add_catalog_skill(
        session,
        canonical_name=competence,
        aliases=[competence],
    )

    skill = add_candidate_skill(
        session,
        candidate_id=CANDIDATE_ID,
        name=competence,
    )

    if avec_preuve:
        add_evidence(
            session,
            candidate_id=CANDIDATE_ID,
            skill_id=skill.id,
            description="Conception de produits digitaux",
        )

    session.add(
        JobOfferDB(
            id=JOB_OFFER_ID,
            title="Product Owner",
            company=company,
            description=f"Nous cherchons du {competence}.",
            status="selected",
        )
    )

    session.commit()

    for evidence in session.query(EvidenceDB).all():
        evidence.experience_id = EXPERIENCE_ID

    session.commit()


def _analyser(required_skills: list[str]):
    from services.matching import analyze_and_save_job_match

    return analyze_and_save_job_match(
        candidate_id=CANDIDATE_ID,
        job_offer_id=JOB_OFFER_ID,
        required_skills=required_skills,
    )


# ============================================================
# HONNETETE DU CONTENU
# ============================================================

def test_la_lettre_affirme_les_competences_prouvees(
    session_factory,
):
    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session)
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    assert lettre.claimed_skills == ["Product Discovery"]
    assert "Product Discovery" in lettre.full_text


def test_la_lettre_n_affirme_jamais_une_competence_sans_preuve(
    session_factory,
):
    """
    Une compétence déclarée sans preuve ne doit pas être présentée
    comme documentée dans la lettre.
    """

    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session, avec_preuve=False)
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    assert lettre.claimed_skills == []
    assert "Product Discovery" in lettre.not_claimed_skills

    # Le nom peut apparaître dans l'objet (intitulé du poste), mais
    # jamais dans une affirmation de compétence.
    affirmations = " ".join(
        paragraphe.text
        for paragraphe in lettre.paragraphs
        if paragraphe.est_une_affirmation
    )

    assert "Product Discovery" not in affirmations


def test_une_competence_deduite_n_est_pas_affirmee(
    session_factory,
):
    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session)

    add_catalog_skill(
        session,
        canonical_name="Python",
        aliases=["Python"],
        skill_id="catalog-python",
    )

    session.close()

    _analyser(["Product Discovery", "Python"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    assert "Python" not in lettre.claimed_skills

    affirmations = " ".join(
        paragraphe.text
        for paragraphe in lettre.paragraphs
        if paragraphe.est_une_affirmation
    )

    assert "Python" not in affirmations


def test_chaque_affirmation_porte_ses_sources(session_factory):
    """
    Un paragraphe qui parle du parcours doit pouvoir être justifié.
    Les paragraphes sans source sont des formules de politesse.
    """

    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session)
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    affirmations = [
        paragraphe
        for paragraphe in lettre.paragraphs
        if paragraphe.est_une_affirmation
    ]

    assert affirmations, "la lettre doit s'appuyer sur le parcours"

    for paragraphe in affirmations:
        assert all(paragraphe.sources)


# ============================================================
# VALIDATION OBLIGATOIRE
# ============================================================

def test_une_lettre_generee_n_est_jamais_validee_d_office(
    session_factory,
):
    """
    Principe de cadrage : aucune lettre n'est finale sans relecture
    de l'utilisateur.
    """

    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session)
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    assert lettre.validated_by_user is False


# ============================================================
# ROBUSTESSE
# ============================================================

def test_le_nom_de_l_entreprise_est_repris_quand_il_existe(
    session_factory,
):
    """
    Pendant du test suivant : sans ce cas, un bug qui perdrait le nom
    de l'entreprise passerait inaperçu, puisque l'absence de nom est
    un comportement valide.
    """

    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session, company="Groupe Meridiem")
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    assert lettre.company == "Groupe Meridiem"
    assert "Groupe Meridiem" in lettre.full_text


def test_une_offre_sans_entreprise_ne_fait_pas_inventer_de_nom(
    session_factory,
):
    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session, company="")
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    assert lettre.company == ""
    assert "votre équipe" in lettre.full_text


def test_la_lettre_reste_correcte_sans_competence_prouvee(
    session_factory,
):
    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session, avec_preuve=False)
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    # La lettre existe, avec ses formules, mais sans paragraphe de
    # compétences fabriqué pour faire nombre.
    assert lettre.salutation in lettre.full_text
    assert lettre.closing in lettre.full_text
    assert "je peux documenter concrètement" not in lettre.full_text


def test_deux_generations_donnent_la_meme_lettre(session_factory):
    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session)
    session.close()

    _analyser(["Product Discovery"])

    jour = date(2026, 9, 2)

    premiere = build_cover_letter(
        CANDIDATE_ID,
        JOB_OFFER_ID,
        redaction_date=jour,
    )

    seconde = build_cover_letter(
        CANDIDATE_ID,
        JOB_OFFER_ID,
        redaction_date=jour,
    )

    assert premiere.full_text == seconde.full_text


def test_la_casse_des_acronymes_est_preservee(session_factory):
    """
    Les éléments du Master CV sont repris tels quels : passer le
    texte en minuscules abîmerait les acronymes (MVP, KPI, IT).
    """

    from database.models import EvidenceDB
    from services.letter import build_cover_letter

    session = session_factory()
    _prepare(session)

    evidence = session.query(EvidenceDB).first()
    evidence.description = "Mise en place d'une logique MVP"
    session.commit()
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    assert "MVP" in lettre.full_text
    assert "mvp" not in lettre.full_text


# ============================================================
# EXPORT
# ============================================================

def test_la_lettre_s_exporte_en_docx_et_pdf(
    session_factory,
    tmp_path,
):
    from services.letter import (
        build_cover_letter,
        export_letter_docx,
        export_letter_pdf,
    )

    session = session_factory()
    _prepare(session)
    session.close()

    _analyser(["Product Discovery"])

    lettre = build_cover_letter(CANDIDATE_ID, JOB_OFFER_ID)

    docx_path = export_letter_docx(lettre, tmp_path / "lettre.docx")
    pdf_path = export_letter_pdf(lettre, tmp_path / "lettre.pdf")

    assert docx_path.exists()
    assert pdf_path.read_bytes().startswith(b"%PDF-")

    from docx import Document

    texte = "\n".join(
        paragraphe.text
        for paragraphe in Document(str(docx_path)).paragraphs
    )

    assert "Madame, Monsieur," in texte
    assert lettre.signature in texte
